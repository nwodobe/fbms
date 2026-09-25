"""Contrôle d'un classeur exporté par l'écran AFLP DATA.

Usage : python3 verifier_classeur.py <fichier.xlsx> [--vide]
Vérifie : nom du fichier, 16 onglets dans l'ordre demandé, en-têtes exacts,
colonnes de contrôle en formules (cash, sacs jute, stock terrain), types des
dates et des nombres, section « Contrôles obligatoires » dans le tableau de bord.
Aucune donnée réelle n'est utilisée : le classeur testé vient de données fictives.
"""
import os, re, sys, datetime
import openpyxl

ENTETES = {
 'AFLP Overview': ['KPI', 'Value', 'Unit', 'Note'],
 'Zones & Clusters': ['Zone', 'Cluster', 'Department', 'Sous-préfecture', 'Main town', 'Zone Head', 'Unit Head', 'Assistant', 'Number of villages', 'Target MT',
   'Potential MT', 'Secured MT', 'Purchased MT', 'Evacuated MT', 'Remaining MT', 'Performance %', 'Risk level', 'Remarks'],
 'Villages Master Data': ['Village ID', 'Village name', 'Zone', 'Cluster', 'Department', 'Sous-préfecture', 'GPS latitude', 'GPS longitude', 'Distance to cluster hub',
   'Road condition', 'Access 10T', 'Access 30T', 'Estimated producers', 'Estimated potential MT', 'Secured potential MT', 'Competition risk', 'Assigned RT', 'Unit Head',
   'Last visit date', 'Village status', 'Remarks'],
 'Producers Registry': ['Producer ID', 'Producer name', 'Phone number', 'Village', 'Zone', 'Cluster', 'RT assigned', 'Estimated farm size', 'Estimated production kg',
   'Secured quantity kg', 'Producer status', 'ID document available', 'Payment method', 'Wave number', 'Last transaction date', 'Total sold kg', 'Total amount paid',
   'Outstanding balance', 'Traceability status', 'Remarks'],
 'RT & Field Teams': ['Staff ID', 'Name', 'Role', 'Zone', 'Cluster', 'Assigned village', 'Phone', 'SIM / Wave number', 'Active status', 'Start date', 'Supervisor',
   'Number of producers assigned', 'Target kg', 'Purchased kg', 'Achievement %', 'Cash advance received', 'Cash justified', 'Cash balance', 'Bags issued', 'Bags returned',
   'Bags balance', 'Incidents', 'Remarks'],
 'Field Missions & Village Visits': ['Mission ID', 'Visit date', 'Zone', 'Cluster', 'Village', 'Staff involved', 'Mission objective', 'Producers met',
   'Estimated volume discussed kg', 'Commitments secured kg', 'Issues raised', 'GPS check-in', 'Photos available', 'Attendance list available', 'Follow-up action',
   'Next visit date', 'Mission status', 'Remarks'],
 'AFLP Daily Purchases': ['Purchase ID', 'Purchase date', 'Zone', 'Cluster', 'Village', 'Producer name', 'Producer code', 'RT name', 'Unit Head', 'Quantity kg', 'Bags count',
   'Price CFA/kg', 'Gross amount CFA', 'Quality observation', 'Moisture if available', 'KOR if available', 'Payment method', 'Payment status', 'Wave transaction ref',
   'Cash voucher ref', 'Purchase status', 'Remarks'],
 'Cash Advances & Payments': ['Transaction ID', 'Date', 'Staff / RT', 'Zone', 'Cluster', 'Village', 'Transaction type', 'Opening advance', 'Amount received',
   'Amount paid to producers', 'Amount returned', 'Difference', 'Current balance', 'Payment method', 'Supporting document', 'Approval status', 'Approved by', 'Remarks',
   'Control (= 0)'],
 'AFLP Jute Bags Ledger': ['Movement ID', 'Movement date', 'Zone', 'Cluster', 'Village', 'RT / Staff', 'Supplier / Producer if applicable', 'Movement type',
   'Bags opening stock', 'Bags received', 'Bags issued to producers', 'Bags used for purchases', 'Bags returned full', 'Bags returned empty', 'Damaged repairable bags',
   'Damaged unusable bags', 'Reconditioned bags', 'Bags transferred out', 'Bags transferred in', 'Closing balance', 'Remarks', 'Control (= 0)'],
 'Field Stock - Village Stock': ['Stock ID', 'Date', 'Zone', 'Cluster', 'Village', 'RT', 'Stock point', 'Opening stock kg', 'Purchases kg', 'Returns kg', 'Evacuated kg',
   'Loss / adjustment kg', 'Closing stock kg', 'Bags in stock', 'Stock status', 'Last physical check', 'Variance kg', 'Remarks', 'Control (= 0)'],
 'Evacuations & Transport': ['Evacuation ID', 'Evacuation date', 'Zone', 'Cluster', 'Origin village / stock point', 'Destination warehouse', 'Truck no', 'Driver name',
   'Transporter', 'Quantity loaded kg', 'Bags loaded', 'Quantity received kg', 'Bags received', 'Difference kg', 'Distance km', 'Transport cost', 'Fuel estimate', 'Status',
   'Incident', 'Remarks'],
 'Warehouse Relay AFLP': ['Warehouse relay', 'Zone', 'Cluster', 'Location', 'Opening stock kg', 'Received from field kg', 'Transferred to Yamoussoukro kg', 'Closing stock kg',
   'Bags received', 'Bags transferred', 'Quality status', 'Lot number', 'Bin / location', 'Last movement date', 'Remarks'],
 'Quality & Traceability': ['Quality ID', 'Date', 'Zone', 'Cluster', 'Village', 'Producer / Lot', 'Sample type', 'Moisture %', 'Nut count', 'KOR', 'Defects', 'Decision',
   'Quality officer', 'Traceability complete', 'Producer linked', 'Village linked', 'RT linked', 'GPS linked', 'Remarks'],
 'Incidents, Risks & Compliance': ['Incident ID', 'Date', 'Zone', 'Cluster', 'Village', 'Reported by', 'Incident type', 'Risk category', 'Description', 'Severity',
   'Immediate action', 'Responsible person', 'Status', 'Closing date', 'Evidence available', 'Remarks'],
 'AFLP Performance Dashboard': ['Section', 'Rank', 'Indicator', 'Value', 'Unit', 'Detail'],
 'AFLP Audit Log': ['Audit ID', 'Date', 'User', 'Role', 'Module', 'Action', 'Object type', 'Object ID', 'Before', 'After', 'Reason', 'Approval', 'Remarks'],
}
FORMULES = {
 'Cash Advances & Payments': ['Opening advance', '+Amount received', '-Amount paid to producers', '-Amount returned', '-Current balance'],
 'AFLP Jute Bags Ledger': ['Bags opening stock', '+Bags received', '+Bags returned full', '+Bags returned empty', '+Bags transferred in', '-Bags issued to producers',
   '-Damaged unusable bags', '-Bags transferred out', '-Closing balance'],
 'Field Stock - Village Stock': ['Opening stock kg', '+Purchases kg', '+Returns kg', '-Evacuated kg', '-Loss / adjustment kg', '-Closing stock kg'],
}
LIBELLES_ABSENCE = ('À compléter', 'Non disponible')


def lettre(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def main(path, vide):
    res = []
    ok = lambda c, t, d='': res.append(('PASS' if c else 'FAIL', t, d))
    ok(re.fullmatch(r'(vide_)?ANAGROCI_AFLP_DATA_\d{4}_\d{8}\.xlsx', os.path.basename(path)) is not None, 'Nom du fichier ANAGROCI_AFLP_DATA_<campagne>_AAAAMMJJ.xlsx', os.path.basename(path))
    wb = openpyxl.load_workbook(path)
    ok(wb.sheetnames == list(ENTETES), '16 onglets dans l’ordre demandé', ' | '.join(wb.sheetnames))
    for name, hdr in ENTETES.items():
        ws = wb[name]
        got = [c.value for c in ws[1]][:len(hdr)]
        ok(got == hdr, f'En-têtes {name} ({len(hdr)} colonnes)', '' if got == hdr else str([(i, a, b) for i, (a, b) in enumerate(zip(got, hdr)) if a != b][:4]))
        ok(ws.freeze_panes == 'A2', f'Ligne d’en-tête figée {name}', str(ws.freeze_panes))
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        if name in ('AFLP Overview', 'AFLP Performance Dashboard'):
            continue
        if vide:
            ok(ws.max_row == 1, f'{name} vide : en-têtes seules', str(ws.max_row))
            continue
        bad = []
        for row in rows:
            for h, v in zip(hdr, row):
                if v is None or v in LIBELLES_ABSENCE or (isinstance(v, str) and v.startswith('=')):
                    continue
                if (h.endswith('date') or h in ('Date', 'Last physical check')) and not isinstance(v, (datetime.datetime, datetime.date)):
                    bad.append((h, v))
                if (h.endswith(' kg') or h.endswith(' MT') or h.startswith('Bags ') or h.startswith('Amount') or h in ('Opening advance', 'Difference', 'Current balance',
                        'Price CFA/kg', 'Gross amount CFA', 'Closing balance', 'Number of villages', 'Incidents')) and not isinstance(v, (int, float)):
                    bad.append((h, v))
        ok(not bad, f'Types {name} (nombres et dates)', str(bad[:4]))
        if name in FORMULES and not vide:
            col = {h: lettre(i + 1) for i, h in enumerate(hdr)}
            for r in range(2, ws.max_row + 1):
                attendu = '=' + ''.join((t[0] if t[0] in '+-' else '') + col[t.lstrip('+-')] + str(r) for t in FORMULES[name])
                f = ws.cell(r, len(hdr)).value
                ok(f == attendu, f'Formule de contrôle {name} ligne {r}', str(f))
    ov = wb['AFLP Overview']
    kpis = [ov.cell(r, 1).value for r in range(2, 26)]
    ok(kpis[0] == 'Campaign' and kpis[1] == 'Target MT' and kpis[23] == 'Last update', 'Overview : 24 KPI dans l’ordre demandé', str(kpis[:3]) + '…' + str(kpis[-1:]))
    ok(ov.cell(3, 2).value == 3000, 'Overview : Target MT = 3 000', str(ov.cell(3, 2).value))
    labels = [ov.cell(r, 1).value for r in range(1, ov.max_row + 1)]
    ok('Filtres appliqués' in labels and 'Notes de lecture' in labels, 'Overview : filtres et notes de lecture')
    pf = wb['AFLP Performance Dashboard']
    col1 = [pf.cell(r, 1).value for r in range(1, pf.max_row + 1)]
    ok('Contrôles obligatoires' in col1, 'Tableau de bord : section « Contrôles obligatoires »')
    if 'Contrôles obligatoires' in col1:
        i = col1.index('Contrôles obligatoires') + 2
        ok([pf.cell(i, c).value for c in range(1, 9)] == ['No', 'Code', 'Control', 'Status', 'Anomalies', 'Value', 'Unit', 'Detail'], 'Tableau de bord : en-têtes des contrôles')
        st = [pf.cell(r, 4).value for r in range(i + 1, pf.max_row + 1)]
        ok(len(st) > 0 and all(s in ('OK', 'ALERTE', 'SANS DONNÉES') for s in st), 'Tableau de bord : statuts des contrôles', str(st))
    for x in res:
        print(f'{x[0]}\t{x[1]}\t{x[2]}')
    fails = [x for x in res if x[0] == 'FAIL']
    print(f"\n{len(res) - len(fails)} PASS · {len(fails)} FAIL")
    return 1 if fails else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], '--vide' in sys.argv))
