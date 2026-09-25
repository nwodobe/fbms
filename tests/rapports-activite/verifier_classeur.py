"""Contrôle d'un classeur exporté par l'écran Rapports d'activité.

Usage : python3 verifier_classeur.py <fichier.xlsx> [--vide]
Vérifie : ordre des onglets, en-têtes identiques au modèle ANAGROCI (en-têtes
seulement, aucune donnée du modèle), types numériques, dates lisibles,
formule de clôture sacherie. Aucune donnée réelle n'est utilisée.
"""
import sys, datetime
import openpyxl

MODELE = {
 'Suppliers': ['No','supplier name','supplier code'],
 'Delivery Plan': ['delivery plan id','planned delivery date','supplier','supplier code','expected quantity kg','expected truck type','expected arrival date','destination warehouse','actual arrival date','delivery status','truck no','remarks'],
 'Truck Reception': ['No','arrival date','offloading date','payment type','warehouse location','warehouse code','supplier name','supplier code','truck no','fiche code','fiche offloading no','bags count','gross weight kg','net weight kg','good bags','humid bags','torn bags','refraction kg','paid weight kg','price cfa per kg','amount cfa','moisture pct','nut count','kor','final kor','origin','remarks','week'],
 'Warehouse Receiving': ['No','warehouse site','warehouse','initial bin no','after drying bin no','lot no','lot dry','delivery note no','scale location','warehouse receipt no','fiche date','activity','ccak validated','storage area','fiche cca no','truck no','arrival date','discharge date','origin','supplier name','cca no','supplier code','good bags offloaded','damaged bags offloaded','total bags discharged','recondition bags','total bags stock','new bags used recondition','good bags','humid bags','torn bags','recondition flag','reconditioned bags','total bags lot','export bags','bio bags','brousse bags','total received bags','gross weight kg','total refraction kg','grn qty kg','paid weight kg','net weight kg','weight difference kg','net weight book kg','grn fresh qty kg','grn dried qty kg','q1 rejection g','q1 void g','q1 oil g','q1 total defect g','q1 good kernel g','q1 immature g','q1 spotted g','q1 total kernels g','q1 kor','q1 murli','q1 shot','q1 fot','q1 useful kernel yield pct','q1 total yield pct','q1 moisture pct','q1 nut count','q1 shell g','q2 rejection g','q2 murli total pct','q2 void g','q2 oil g','q2 total defect g','q2 good kernel g','q2 immature g','q2 spotted g','q2 total kernels g','q2 kor','q2 murli','q2 after dry kor','q2 shot','q2 fot','q2 useful kernel yield pct','q2 total yield pct','q2 moisture pct','q2 nut count','q2 shell g'],
 'Quality Inspection': ['quality inspection id','reception id','inspection stage','warehouse','activity type','inspection date','supplier name','ccak code','anagroci code','truck no','lot no','origin','moisture pct','nut count','good kernel g','spotted g','immature g','void g','oil g','browns rejection g','kor','decision','quality head','remarks'],
 'Drying Batch': ['No','drying date','warehouse site','warehouse','source bin no','destination bin no','lot no raw','warehouse receipt no raw','fiche no raw','origin raw','supplier name raw','supplier code raw','ccak supplier code raw','status','dry type','destination raw','issued gross with bags pallets kg','issued pallet weight kg','issued gross with bags kg','issued bags','issued net weight kg','received gross with bags pallets kg','received pallet weight kg','received gross with bags kg','received bags','received net weight kg','input moisture pct','input nut count','input kor','output moisture pct','output nut count','output kor','oil','rejection','good kernel','immature','spotted','shell','void','moisture loss pct','drying loss kg','drying loss pct','triage loss kg','triage loss pct','damaged nuts kg','input bags for drying','output bags after drying','drying or picking','remarks','issued to production','re drying date','moisture first redry pct','moisture second dry pct','issued to production 3','remarks 4','needs lot allocation clarification'],
 'Warehouse Activity Ledger': ['activity date','activity type','particulars','offloaded','bags','quantity mt','fiche no','lot no','vendor','rate','total','reference no','remarks','gross weight with pallet kg','pallet count','gross weight kg','gross weight without pallet kg','net weight kg'],
 'Jute Bags Movement': ['No','movement date','activity type','supplier code','warehouse location','truck no','warehouse receipt no','bags issued','bags received'],
}
ORDRE = list(MODELE.keys()) + ['Jute Bags Balance', 'Summary']
NUM = ('kg','bags','count','pct','kor','qty','mt','g','cfa','opening','receipts','transfers','issues','closing','variance','damaged','No')
DATE = ('date',)

def main(path, vide):
    wb = openpyxl.load_workbook(path)  # formules conservées
    res = []
    ok = lambda c, t, d='': res.append(('PASS' if c else 'FAIL', t, d))
    ok(wb.sheetnames == ORDRE, "Ordre des onglets", ' | '.join(wb.sheetnames))
    for name, hdr in MODELE.items():
        ws = wb[name]
        got = [c.value for c in ws[1]][:len(hdr)]
        ok(got == hdr, f"En-têtes {name} ({len(hdr)} colonnes)", '' if got == hdr else f"écart : {[ (i,a,b) for i,(a,b) in enumerate(zip(got,hdr)) if a!=b][:5]}")
        ok(ws.max_column == len(hdr), f"Pas de colonne en trop {name}", str(ws.max_column))
        if vide:
            ok(ws.max_row == 1, f"{name} vide : en-têtes seules", str(ws.max_row))
            continue
        bad = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            for h, v in zip(hdr, row):
                if v is None: continue
                last = h.split()[-1]
                if any(h.endswith(s) or last == s for s in DATE) and not isinstance(v, (datetime.datetime, datetime.date)):
                    bad.append((h, v))
                elif (last in NUM or h in ('No','rate','total','oil','rejection','good kernel','immature','spotted','shell','void')) and not isinstance(v, (int, float)):
                    bad.append((h, v))
        ok(not bad, f"Types {name} (nombres et dates)", str(bad[:4]))
    b = wb['Jute Bags Balance']
    hb = [c.value for c in b[1]]
    ok(hb[:11] == ['warehouse location','warehouse code','opening','receipts','transfers in','issues','damaged/discarded','transfers out','closing','closing actual (ledger)','variance'], "En-têtes Jute Bags Balance", str(hb))
    for r in range(2, b.max_row + 1):
        f = b.cell(r, 9).value; fv = b.cell(r, 11).value
        ok(isinstance(f, str) and f == f"=C{r}+D{r}+E{r}-F{r}-G{r}-H{r}", f"Formule Closing ligne {r}", str(f))
        ok(isinstance(fv, str) and fv == f"=J{r}-I{r}", f"Formule Variance ligne {r}", str(fv))
        o, rc, ti, iss, dmg, to = (b.cell(r, c).value for c in range(3, 9))
        ok(all(isinstance(x, (int, float)) for x in (o, rc, ti, iss, dmg, to)), f"Valeurs numériques balance ligne {r}")
    s = wb['Summary']
    ok(s['A1'].value == 'Indicateur' and s.max_row > 5, "Onglet Summary renseigné", str(s.max_row))
    fails = [x for x in res if x[0] == 'FAIL']
    for x in res: print(f"{x[0]}\t{x[1]}\t{x[2]}")
    print(f"\n{len(res)-len(fails)}/{len(res)} contrôles conformes")
    return 1 if fails else 0

if __name__ == '__main__':
    sys.exit(main(sys.argv[1], '--vide' in sys.argv))
